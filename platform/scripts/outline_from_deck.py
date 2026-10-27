"""Write slides/outline.md from the owner's final deck: every slide's title, text, links and notes.

    uv run python platform/scripts/outline_from_deck.py                # the default deck
    uv run python platform/scripts/outline_from_deck.py --deck PATH --out slides/outline.md

The final deck (DECISIONS.md, 28 October 2026) is the owner's own file, edited by hand and final:
this script only reads it, through a temporary copy, and never writes it. The outline is a
verbatim transcript, so the deck's own words sit in fenced text blocks and each title in a code
span: the course's naming and placeholder rules read the course's prose, not the owner's slides.
Three things are not copied word for word: an email address becomes {email address}, so a
private contact never reaches a published file; an emoji icon becomes {icon}, as the course's
prose lint asks of every published page; and a link is listed by its address. The header records
the deck's slide count and SHA-256, which `check_outline.py` reads where the deck is absent.

The deck is found at --deck, else $LOOP_ENGINEERING_DECK, else the owner's Desktop copy, else the
backup in slides/backup/ (both kept out of the release).
"""

from __future__ import annotations

import argparse
import hashlib
import os
import re
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lint_prose import _EMOJI, ALLOWED_SYMBOLS  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
OUTLINE = ROOT / "slides" / "outline.md"
DECK_NAME = "Loop Engineering - Slides.pptx"
CANDIDATES = (
    Path.home() / "Desktop" / DECK_NAME,
    ROOT / "slides" / "backup" / "Loop Engineering - Slides (final, 28 October 2026).pptx",
)
EMAIL = re.compile(r"(mailto:)?[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
FINGERPRINT = re.compile(r"^Deck: (\d+) slides, SHA-256 ([0-9a-f]{64})\.$", re.M)


def find_deck(given: str | None = None) -> Path | None:
    for candidate in (given, os.environ.get("LOOP_ENGINEERING_DECK"), *CANDIDATES):
        if candidate and Path(candidate).is_file():
            return Path(candidate)
    return None


def _icon(match: re.Match[str]) -> str:
    found = match.group(0)
    if found in ALLOWED_SYMBOLS:
        return found
    return "" if found == "️" else "{icon}"  # the variation selector shows nothing


def _private(text: str) -> str:
    return _EMOJI.sub(_icon, EMAIL.sub("{email address}", text))


def _fence(text: str) -> str:
    """A fenced text block that no line of `text` can close early."""
    longest = max((len(m.group(0)) for m in re.finditer(r"`{3,}", text)), default=0)
    fence = "`" * max(3, longest + 1)
    return f"{fence}text\n{text.rstrip()}\n{fence}"


def _texts(shapes, lines: list[str], links: list[str]) -> None:
    for shape in shapes:
        if shape.shape_type == 6:  # a group
            _texts(shape.shapes, lines, links)
            continue
        frames = []
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [c.text_frame.text.replace("\n", " ").strip() for c in row.cells]
                lines.append("| " + " | ".join(cells) + " |")
                frames += [c.text_frame for c in row.cells]
        elif shape.has_text_frame:
            frames.append(shape.text_frame)
            for paragraph in shape.text_frame.paragraphs:
                line = "".join(run.text for run in paragraph.runs).strip()
                if line:
                    lines.append(line)
        for frame in frames:
            for paragraph in frame.paragraphs:
                for run in paragraph.runs:
                    if run.hyperlink and run.hyperlink.address:
                        links.append(run.hyperlink.address)
        try:
            address = shape.click_action.hyperlink.address
        except Exception:  # a shape type with no click action
            address = None
        if address:
            links.append(address)


def outline(deck: Path) -> str:
    from pptx import Presentation

    digest = hashlib.sha256(deck.read_bytes()).hexdigest()
    with tempfile.TemporaryDirectory() as tmp:
        copy = Path(tmp) / "deck.pptx"
        shutil.copyfile(deck, copy)
        slides = list(Presentation(str(copy)).slides)
        out = [
            "# Loop Engineering: deck outline",
            "",
            "The class deck as the owner finished it, slide by slide: the title, the text on "
            "the slide, its links and its speaker notes. It is generated from the final deck by "
            "`platform/scripts/outline_from_deck.py` and never edited by hand. The deck's words "
            "are kept as written, in text blocks; an email address shows as {email address} and "
            "an emoji icon as {icon}. Terms the slides use include Structured Query Language "
            "(SQL), application programming interface (API) and large language model (LLM).",
            "",
            f"Deck: {len(slides)} slides, SHA-256 {digest}.",
        ]
        for number, slide in enumerate(slides, 1):
            title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
            lines: list[str] = []
            links: list[str] = []
            _texts(slide.shapes, lines, links)
            notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
            heading = f"## Slide {number}"
            if title:
                heading += ": `" + _private(" ".join(title.split())).replace("`", "'") + "`"
            shown = _private("\n".join(lines)) or "(no text)"
            out += ["", heading, "", "On the slide:", "", _fence(shown)]
            unique = list(dict.fromkeys(_private(a) for a in links))
            if unique:
                out += ["", "Links:", ""]
                web = ("http://", "https://")
                out += [f"- <{a}>" if a.startswith(web) else f"- `{a}`" for a in unique]
            out += ["", "Speaker notes:", "", _fence(_private(notes) or "(no notes)")]
        return "\n".join(out) + "\n"


def recorded_fingerprint(text: str) -> tuple[int, str] | None:
    found = FINGERPRINT.search(text)
    return (int(found.group(1)), found.group(2)) if found else None


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--deck", help="the final deck (read only)")
    ap.add_argument("--out", type=Path, default=OUTLINE)
    a = ap.parse_args(argv)
    deck = find_deck(a.deck)
    if deck is None:
        print("x the final deck is not here: pass --deck or set LOOP_ENGINEERING_DECK")
        return 2
    a.out.write_text(outline(deck), encoding="utf-8", newline="\n")
    print(f"wrote {a.out} from {deck.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
