"""The outline check and its generator: slides/outline.md is a transcript of the owner's final deck.

The owner's decision of 28 October 2026 (DECISIONS.md) made the final deck the outline's source in
place of slides/source, so these tests build a small deck of their own and never read the owner's.
"""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pytest

SCRIPT = Path(__file__).resolve().parents[2] / "scripts" / "check_outline.py"


def _load():
    spec = importlib.util.spec_from_file_location("check_outline", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    sys.modules["check_outline"] = module
    spec.loader.exec_module(module)
    return module


check_outline = _load()
generator = check_outline.GENERATOR


def _deck(path: Path, notes: str = "Timing: 2 minutes.") -> Path:
    from pptx import Presentation

    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[1])
    first.shapes.title.text = "The agent loop"
    body = first.placeholders[1].text_frame
    body.text = "Write to the course team at someone@example.com"
    run = body.add_paragraph().add_run()
    run.text = "The reference"
    run.hyperlink.address = "https://example.com/paper"
    first.notes_slide.notes_text_frame.text = notes
    deck.slides.add_slide(deck.slide_layouts[6])  # a blank slide: no title, no text, no notes
    deck.save(str(path))
    return path


def test_the_outline_transcribes_every_slide(tmp_path: Path) -> None:
    deck = _deck(tmp_path / "deck.pptx")
    before = deck.read_bytes()
    text = generator.outline(deck)
    assert deck.read_bytes() == before, "the deck is read only"
    assert generator.recorded_fingerprint(text)[0] == 2
    assert "## Slide 1: `The agent loop`" in text
    assert "## Slide 2\n" in text
    assert "- <https://example.com/paper>" in text
    assert "Timing: 2 minutes." in text
    assert "(no text)" in text and "(no notes)" in text


def test_an_email_address_never_reaches_the_outline(tmp_path: Path) -> None:
    text = generator.outline(_deck(tmp_path / "deck.pptx", notes="Ask mailto:a.b@example.org"))
    assert "example.org" not in text and "someone@example.com" not in text
    slides = text.split("\n## Slide 1", 1)[1]  # the header names the placeholder itself
    assert slides.count("{email address}") == 2


def test_an_emoji_icon_shows_as_a_placeholder(tmp_path: Path) -> None:
    text = generator.outline(_deck(tmp_path / "deck.pptx", notes="Chat \U0001f4ac, done ✓"))
    assert "\U0001f4ac" not in text
    assert "Chat {icon}, done ✓" in text  # the plain status symbols the prose lint allows stay


def test_a_fence_in_the_notes_cannot_close_the_block_early() -> None:
    assert generator._fence("a\n```\nb").startswith("````text\n")


def test_a_whole_outline_passes_on_its_own(tmp_path: Path) -> None:
    assert check_outline.whole_problems(generator.outline(_deck(tmp_path / "deck.pptx"))) == []


def test_an_outline_without_a_fingerprint_fails() -> None:
    found = check_outline.whole_problems("# Outline\n\n## Slide 1\n")
    assert found == [
        "slides/outline.md records no deck fingerprint: run make outline on the final deck"
    ]


def test_an_outline_missing_a_slide_fails(tmp_path: Path) -> None:
    text = generator.outline(_deck(tmp_path / "deck.pptx")).replace("## Slide 2\n", "")
    assert check_outline.whole_problems(text)[0].startswith("slides/outline.md records 2 slides")


def test_a_stale_outline_fails_and_shows_the_difference() -> None:
    found = check_outline.match_problems("# Outline\nnew line\n", "# Outline\nold line\n")
    assert len(found) == 1
    assert found[0].startswith("slides/outline.md does not match the final deck")
    assert "-old line" in found[0] and "+new line" in found[0]


def _tree(tmp_path: Path, outline: str | None) -> Path:
    root = tmp_path / "tree"
    (root / "slides").mkdir(parents=True)
    if outline is not None:
        (root / "slides" / "outline.md").write_bytes(outline.encode("utf-8"))
    return root


@pytest.mark.parametrize(
    ("change", "code"),
    [
        (lambda t: t, 0),
        (lambda t: t.replace("\n", "\r\n"), 1),  # byte for byte: line endings count
        (lambda t: t.replace("The agent loop", "The event loop"), 1),
    ],
)
def test_main_compares_the_outline_with_the_deck(tmp_path: Path, change, code: int) -> None:
    deck = _deck(tmp_path / "deck.pptx")
    root = _tree(tmp_path, change(generator.outline(deck)))
    assert check_outline.main(["--root", str(root), "--deck", str(deck)]) == code


def test_main_checks_the_outline_on_its_own_where_the_deck_is_absent(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    monkeypatch.setattr(generator, "find_deck", lambda given=None: None)
    text = generator.outline(_deck(tmp_path / "deck.pptx"))
    assert check_outline.main(["--root", str(_tree(tmp_path, text))]) == 0
    assert "the final deck is not here" in capsys.readouterr().out


def test_a_missing_outline_fails(tmp_path: Path) -> None:
    assert check_outline.main(["--root", str(_tree(tmp_path, None))]) == 1
