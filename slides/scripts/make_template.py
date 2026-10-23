"""Make a slide template from a finished deck: keep its design, drop its content.

    uv run --with python-pptx python slides/scripts/make_template.py \
        --deck /tmp/deck-copy.pptx --out slides/template.pptx

Always pass a COPY of the deck: the script never writes to --deck, but a deck someone is editing
should not be opened by a script at all. What it does, in order:

1. removes every slide (and with them their pictures, charts and notes);
2. keeps the slide masters, the layouts, the theme and the notes master;
3. blanks any text a master or layout carries in a placeholder (course titles, footers);
4. drops the thumbnail, the comment authors, and the document properties that name a course
   or a person;
5. adds the example slides of `EXAMPLES` (title, section divider, concept, diagram, demo, result
   table, notes example), each with speaker notes in the course's notes convention.

The example slides use the layouts by name (`TITLE_ONLY`, `OBJECT`, ...) and fall back to the
first layout with a title placeholder when a name is missing, so the script works on other decks.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Pt

NEUTRAL_TITLE = "<COURSE>: slide template"

# (layout name, title, subtitle or body lines, extra, notes)
EXAMPLES: list[dict] = [
    {
        "kind": "title",
        "layout": ("TITLE_4", "TITLE", "TITLE_1"),
        "title": "<COURSE TITLE>",
        "body": ["<one line: what the learner builds, then how it is measured>"],
        "notes": [
            "Time: 0.5 min (0:00-0:00:30)",
            "Say the course title and the one question the class answers.",
            "Point at: the subtitle; it names the build and the measurement.",
            'Transition: "First, who is teaching and how the class runs."',
        ],
    },
    {
        "kind": "section",
        "layout": ("SECTION_HEADER_2_1", "SECTION_HEADER", "SECTION_HEADER_2"),
        "title": "Section <N>: <SECTION NAME>",
        "body": ["<the failure this section fixes, in one line>"],
        "notes": [
            "Time: 0.5 min (<start>-<end>)",
            "Name the section and the failure it answers; no new terms yet.",
            "Before this slide: launch the section's demo in terminal 2 (make show-<N>), so it "
            "has finished by the demo slide.",
            'Transition: "Here is what goes wrong without it."',
        ],
    },
    {
        "kind": "concept",
        "layout": ("TITLE_ONLY",),
        "title": "<N>) <Concept name>: <what it does, in a few words>",
        "body": [
            "<One sentence: what the concept is.>",
            "<Bullet: what the agent sees or keeps because of it>",
            "<Bullet: what checks it, and where (file path in the repo)>",
            "<Bullet: what it costs>",
        ],
        "key": "Key idea:  <the one sentence a learner should repeat>",
        "notes": [
            "Time: 4 min (<start>-<end>)",
            'Source <n>: <title> (<link>): "<short quote that the slide claims>".',
            "Point at: the key idea box last; read it aloud once.",
            "In the repo: <section folder>/<file> holds this; open it only if asked.",
            "Re-check after live runs: any count on this slide against docs/evidence/.",
        ],
    },
    {
        "kind": "diagram",
        "layout": ("TITLE_ONLY",),
        "title": "<Diagram title: the mechanism in one picture>",
        "body": ["<One line: what to read first in the picture>"],
        "box": "DIAGRAM: <section folder>/diagrams/<name>.png\n(exported from <name>.excalidraw)",
        "notes": [
            "Time: 3 min (<start>-<end>)",
            "Walk the picture left to right; name each box once, in plain words.",
            "Point at: <the box or arrow that carries the idea>.",
            "The diagram file is the source of truth; if it changes, re-export and replace.",
        ],
    },
    {
        "kind": "demo",
        "layout": ("OBJECT", "TITLE_AND_BODY_4", "Title and Content"),
        "title": "Section <N>: <SECTION NAME> - Implementation Notes: <demo title>",
        "body": ["DEMO - <N>"],
        "box": "console screen <N>\n(a screenshot from the demo guide, or the live console)",
        "notes": [
            "Time: 10 min (<start>-<end>)",
            "Files for this demo: <section folder>/demo/run.sh (the command), "
            "<section folder>/demo/walkthrough.md (the code, explained), "
            "<section folder>/demo/expected/ (the captured run, the fallback).",
            "Switch to browser tab 2, the run launched at the section start.",
            "Screen <N> (1): <what to point at first>. (2): <second>. (3): <the result>.",
            "If the live run failed or is late: open the captured run in demo/expected/ and say "
            "so.",
            "Re-check after live runs: the counts you read aloud against the run on screen.",
        ],
    },
    {
        "kind": "table",
        "layout": ("TITLE_ONLY",),
        "title": "Results: <what was measured>, one row per <unit>",
        "body": ["<One line: how to read the table; counts, not percentages, for small n.>"],
        "table": [
            ["Section", "Without the layer", "With the layer", "Difference"],
            ["<Section 1 name>", "<k of n>", "<k of n>", "<none measurable | +k>"],
            ["<Section 2 name>", "<k of n>", "<k of n>", "<none measurable | +k>"],
            ["<Section 3 name>", "<k of n>", "<k of n>", "<none measurable | +k>"],
        ],
        "notes": [
            "Time: 3 min (<start>-<end>)",
            "Read one row aloud, then the difference column.",
            "Every number on this slide comes from docs/evidence/<file>.md; quote the run there, "
            "not from memory.",
            "Re-check after live runs: today's runs add one row or one dot; update before class.",
        ],
    },
    {
        "kind": "notes",
        "layout": ("TITLE_ONLY",),
        "title": "What to say: the speaker-note convention",
        "body": [
            "Line 1: Time: <minutes> min (<start>-<end>), from the class start.",
            "Then: what to say, in the order to say it; quotes with their source number.",
            "Point at: what the room should look at, slide or screen.",
            "Transition: the sentence that leads to the next slide.",
            "Re-check after live runs: every number or claim that a live run can change.",
        ],
        "notes": [
            "Time: 1 min (<start>-<end>)",
            'Say: "<the sentence you open with>."',
            "Point at: the first line of the slide.",
            'Likely question: "<question>" Answer: <one line, and where in the repo it is>.',
            "Re-check after live runs: nothing on this slide.",
            'Transition: "<the sentence that leads to the next slide>."',
        ],
    },
]


def find_layout(prs: Presentation, names: tuple[str, ...]):
    """The first layout (of any master) with one of `names`, else one with a title."""
    layouts = [lay for master in prs.slide_masters for lay in master.slide_layouts]
    for name in names:
        for lay in layouts:
            if lay.name == name:
                return lay
    for lay in layouts:
        if any(ph.placeholder_format.type in (1, 3) for ph in lay.placeholders):
            return lay
    return layouts[0]


def remove_slides(prs: Presentation) -> int:
    sld_ids = prs.slides._sldIdLst
    count = 0
    for sld_id in list(sld_ids):
        prs.part.drop_rel(sld_id.rId)
        sld_ids.remove(sld_id)
        count += 1
    return count


def blank_layout_text(prs: Presentation) -> int:
    """Empty the text of master and layout placeholders, keeping their formatting; the slide
    number field (‹#›) stays."""
    blanked = 0
    holders = []
    for master in prs.slide_masters:
        holders += list(master.placeholders)
        for lay in master.slide_layouts:
            holders += list(lay.placeholders)
    for ph in holders:
        if not ph.has_text_frame:
            continue
        for para in ph.text_frame.paragraphs:
            for run in para.runs:
                if run.text.strip():
                    run.text = ""
                    blanked += 1
    return blanked


def clean_package(prs: Presentation) -> None:
    """No thumbnail of the old first slide, no comment authors, neutral properties."""
    pkg = prs.part.package
    for rel_id, rel in list(pkg._rels.items()):
        if rel.reltype.endswith("/thumbnail"):
            pkg._rels.pop(rel_id)
    for rel_id, rel in list(prs.part.rels.items()):
        if rel.reltype.endswith("/commentAuthors"):
            prs.part.drop_rel(rel_id)
    props = prs.core_properties
    props.title = NEUTRAL_TITLE
    props.subject = ""
    props.author = "<AUTHOR>"
    props.last_modified_by = "<AUTHOR>"
    props.keywords = ""
    props.comments = ""
    props.category = ""


def _textbox(slide, x, y, w, h, lines, size=20, bold_first=False):
    box = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = line
        for run in para.runs:
            run.font.size = Pt(size)
            run.font.bold = bold_first and i == 0
    return box


def add_example(prs: Presentation, spec: dict) -> None:
    slide = prs.slides.add_slide(find_layout(prs, spec["layout"]))
    W, H = prs.slide_width, prs.slide_height
    holders = {ph.placeholder_format.type: ph for ph in slide.placeholders}
    title = holders.get(1) or holders.get(3)
    if title is not None:
        title.text = spec["title"]
    else:
        _textbox(slide, W * 0.05, H * 0.05, W * 0.9, H * 0.12, [spec["title"]], 28)
    body = holders.get(4) or holders.get(2)
    top = int(H * 0.2)
    if spec["kind"] in ("title", "section", "demo") and body is not None:
        body.text = "\n".join(spec["body"])
    elif spec["kind"] == "concept":
        _textbox(slide, W * 0.05, top, W * 0.9, H * 0.08, spec["body"][:1], 18)
        _textbox(slide, W * 0.07, top + H * 0.1, W * 0.86, H * 0.4, spec["body"][1:], 18)
        _textbox(slide, W * 0.05, H * 0.75, W * 0.9, H * 0.1, [spec["key"]], 18, bold_first=True)
    else:
        _textbox(slide, W * 0.05, top, W * 0.9, H * 0.08, spec["body"], 18)
    if "box" in spec:
        shape = slide.shapes.add_shape(1, Emu(W * 0.15), Emu(H * 0.42), Emu(W * 0.7), Emu(H * 0.4))
        shape.text_frame.text = spec["box"]
        shape.fill.background()
    if "table" in spec:
        rows = spec["table"]
        frame = slide.shapes.add_table(
            len(rows), len(rows[0]), Emu(W * 0.05), Emu(H * 0.32), Emu(W * 0.9), Emu(H * 0.4)
        )
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                frame.table.cell(r, c).text = value
    slide.notes_slide.notes_text_frame.text = "\n".join(spec["notes"])


def make_template(deck: Path, out: Path) -> dict:
    if deck.resolve() == out.resolve():
        raise ValueError("--out must differ from --deck: the deck is never written")
    prs = Presentation(str(deck))
    removed = remove_slides(prs)
    blanked = blank_layout_text(prs)
    clean_package(prs)
    for spec in EXAMPLES:
        add_example(prs, spec)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    check = Presentation(str(out))  # it must load again
    layouts = sum(len(m.slide_layouts) for m in check.slide_masters)
    return {
        "removed": removed,
        "blanked": blanked,
        "masters": len(check.slide_masters),
        "layouts": layouts,
        "slides": len(check.slides),
        "bytes": out.stat().st_size,
    }


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--deck", type=Path, required=True, help="a COPY of the finished deck")
    ap.add_argument("--out", type=Path, required=True)
    a = ap.parse_args(argv)
    info = make_template(a.deck, a.out)
    print(
        f"{a.out.name}: removed {info['removed']} slides, blanked {info['blanked']} layout texts; "
        f"{info['masters']} masters, {info['layouts']} layouts, {info['slides']} example slides, "
        f"{info['bytes'] / 1e6:.2f} MB"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
