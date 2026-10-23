"""Extract the text and speaker notes of a .pptx into Markdown, for a read-only slide review.

    uv run python slides/scripts/extract_deck.py DECK.pptx [--out deck.md]

One `## Slide N: <title>` block per slide, with the layout name, every text frame (grouped
shapes and table cells included, top to bottom, then left to right), and the notes as a quote.
It only reads the deck; copy a deck that someone is editing before you point this at it.
"""

from __future__ import annotations

import argparse
import sys
from collections.abc import Iterable, Iterator
from pathlib import Path

from pptx import Presentation


def _shape_texts(shapes: Iterable) -> Iterator[str]:
    ordered = sorted(shapes, key=lambda s: ((s.top or 0) // 50000, s.left or 0))
    for shape in ordered:
        if shape.shape_type == 6:  # group
            yield from _shape_texts(shape.shapes)
            continue
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                yield "| " + " | ".join(c.text.replace("\n", " ").strip() for c in row.cells) + " |"
            continue
        if shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text and text != "‹#›":
                yield text


def _title(slide) -> str:
    title = slide.shapes.title
    if title is not None and title.text_frame.text.strip():
        return " ".join(title.text_frame.text.split())
    for text in _shape_texts(slide.shapes):
        return " ".join(text.split())[:80]
    return "(no title)"


def deck_markdown(path: Path) -> str:
    prs = Presentation(str(path))
    out = [f"# {path.name}: text and notes", "", f"{len(prs.slides)} slides.", ""]
    for n, slide in enumerate(prs.slides, start=1):
        out.append(f"## Slide {n}: {_title(slide)}")
        out.append("")
        out.append(f"Layout: `{slide.slide_layout.name}`")
        out.append("")
        for text in _shape_texts(slide.shapes):
            for line in text.splitlines():
                out.append(line if line.startswith("| ") else f"- {line.strip()}")
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        out.append("")
        out.append("Notes:")
        out.append("")
        out += [f"> {line}" for line in notes.splitlines()] if notes else ["> (none)"]
        out.append("")
    return "\n".join(out)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("deck", type=Path)
    ap.add_argument("--out", type=Path, default=None)
    a = ap.parse_args(argv)
    text = deck_markdown(a.deck)
    if a.out:
        a.out.write_text(text + "\n", encoding="utf-8")
        print(f"wrote {a.out.name}")
    else:
        sys.stdout.write(text + "\n")
    return 0


if __name__ == "__main__":
    sys.exit(main())
